#!/usr/bin/env python3
"""
export_to_docx.py – Gera arquivo DOCX com os artigos aprovados da Revista Ateísta.

Uso:
    python export_to_docx.py --data-dir dados/ --out-file revista.docx

Requer: pandas, python-docx, bbcode, bleach.
"""
from __future__ import annotations

import argparse
import ast
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
import bbcode
import bleach

BB = bbcode.Parser()

def bb2text(txt: str) -> str:
    """Converte BBCode básico em texto simples."""
    html = BB.format(txt or "")
    html = html.replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
    return bleach.clean(html, tags=[], strip=True)

def clean_id(v: Any) -> str | None:
    """Normaliza IDs que podem vir como lista, string ou objeto."""
    if v is None:
        return None
    if isinstance(v, str) and v.startswith("["):
        try:
            lst = ast.literal_eval(v)
            if isinstance(lst, list) and lst:
                return clean_id(lst[0])
        except Exception:
            return v
    if isinstance(v, dict):
        return v.get("unique_id") or v.get("_id")
    return str(v)

def latest_csv(folder: Path, prefix: str) -> Path:
    """Retorna o CSV mais recente com determinado prefixo."""
    return sorted(folder.glob(f"{prefix}_*.csv"))[-1]

def build_docx(articles: list[dict], edition: dict, out_file: Path) -> None:
    """Monta o documento DOCX com título da edição e artigos."""
    doc = Document()
    doc.add_heading(edition.get("titulo") or f"Edição nº {edition.get('numero', '')}", 0)
    for art in articles:
        doc.add_heading(art.get("titulo", "Sem título"), level=1)
        if art.get("autor_nome"):
            doc.add_paragraph(art["autor_nome"])
        doc.add_paragraph(bb2text(art.get("conteudo")))
    doc.save(out_file)

def main() -> None:
    ap = argparse.ArgumentParser(description="Exporta artigos em DOCX")
    ap.add_argument("--data-dir", type=Path, required=True, help="Diretório com CSVs")
    ap.add_argument("--out-file", type=Path, default=Path("revista_ateista.docx"))
    args = ap.parse_args()

    ed_df = pd.read_csv(latest_csv(args.data_dir, "edicoes"))
    art_df = pd.read_csv(latest_csv(args.data_dir, "artigos_enriched"))

    latest_num = ed_df["numero"].max()
    edition = ed_df[ed_df["numero"] == latest_num].iloc[0].to_dict()
    ed_id = edition["_id"]

    mask = (art_df["status"] == "Aprovado") & (art_df["edicao"].apply(clean_id) == ed_id)
    articles = art_df[mask].to_dict(orient="records")

    build_docx(articles, edition, args.out_file)
    print("✅ DOCX gerado:", args.out_file)

if __name__ == "__main__":
    main()
